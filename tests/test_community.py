from copy import deepcopy
from pathlib import Path
import json
import os
import subprocess
import sys
import pytest
from fintube_hx.engine import run_case,capabilities
from fintube_hx.schema import validate_case,EditionCapabilityError
from fintube_hx.scan import scan_case
from fintube_hx.spreadsheet import export_workbook

ROOT=Path(__file__).resolve().parents[1]
REFERENCE=json.loads((ROOT/'tests/baseline-reference.json').read_text(encoding='utf8'))


def example(name='water_dry.json'):
    return json.loads((ROOT/'examples'/name).read_text(encoding='utf8'))


@pytest.mark.parametrize('filename',sorted(REFERENCE))
def test_approved_baselines_agree_with_historical_record(filename):
    c=example(filename);snapshot=deepcopy(c);r=run_case(c);ref=REFERENCE[filename]
    assert c==snapshot and r['converged']
    assert r['heat_transfer_w']==pytest.approx(ref['heat_transfer_w'],rel=1e-7)
    assert r['fluid_out']['temperature_c']==pytest.approx(ref['fluid_out_t_c'],abs=1e-7)
    assert r['energy_balance_relative_error']<=.005
    assert r['edition']=='community' and r['provenance']['software']=='fintube-hx'
    if 'target' in c:
        assert r['target_solution']['cold_verified']
        assert abs(r['target_solution']['residual_k'])<=c['target']['tolerance_k']
        assert r['target_solution']['mass_flow_kg_s']==pytest.approx(ref['mass_flow_kg_s'],rel=1e-7)
    assert 'tube_results' not in r and 'air_column_results' not in r


@pytest.mark.parametrize('field',['network','velocity_field','water_model','refrigerant_model','air_arrangement','capacity'])
def test_private_modes_explicitly_rejected(field):
    c=example()
    if field=='network':c['network']={'paths':[[1]]}
    elif field=='velocity_field':c['air']['velocity_field']={'values':[[1]]}
    elif field=='air_arrangement':c['model'][field]='column_stream_crossflow'
    elif field in ('water_model','refrigerant_model'):c['model'][field]='tube_network'
    else:
        c=example('condenser.json');c['target']={'heat_transfer_w':1000,'condensing_temperature_bounds_c':[40,50]}
    with pytest.raises(EditionCapabilityError):run_case(c)


@pytest.mark.parametrize('field',['axial_segments_per_tube','heat_transfer_multiplier','return_bend_equivalent_length_m','refrigerant_friction_multiplier','typo'])
def test_unrecognized_active_fields_not_silently_ignored(field):
    c=example();c['model'][field]=1
    with pytest.raises(ValueError,match='Unsupported model fields'):validate_case(c)


@pytest.mark.parametrize('change',['two_inlets','wrong_kind_key','rh_percent','hidden_ua','inactive_geometry','nonfinite_metadata'])
def test_invalid_baseline_inputs(change):
    c=example()
    if change=='two_inlets':c['fluid']['inlet_h_j_kg']=1
    elif change=='wrong_kind_key':c['model']['refrigerant_model']='lumped'
    elif change=='rh_percent':c['air']['relative_humidity']=50
    elif change=='hidden_ua':c=example('water_geometry.json');c['model']['ua_w_k']=100
    elif change=='inactive_geometry':c['geometry']={'rows':2}
    else:c.setdefault('metadata',{})['x']=float('inf')
    with pytest.raises(ValueError):run_case(c)


def test_scan_retains_failed_points():
    r=scan_case(example(),'air.relative_humidity',[.2,50])
    assert r['passed_points']==1 and r['failed_points']==1 and 'error' in r['results'][1]


def test_native_excel_preserves_literal_text_and_empty_pressure_scope(tmp_path):
    from openpyxl import load_workbook
    c=example();c['name']='=1+1';r=run_case(c);p=tmp_path/'result.xlsx';export_workbook(p,[r])
    wb=load_workbook(p);assert '逐管结果' not in wb.sheetnames
    assert wb['汇总']['A2'].data_type=='s' and wb['汇总']['A2'].value=='=1+1'
    keys=[c.value for c in wb['汇总'][1]];col=keys.index('water_straight_tube_pressure_drop_pa')+1
    assert wb['汇总'].cell(2,col).value is None;wb.close()


def test_public_cli_has_only_public_commands(tmp_path):
    env=dict(os.environ,PYTHONPATH=str(ROOT/'src'),PYTHONIOENCODING='utf-8')
    p=subprocess.run([sys.executable,'-m','fintube_hx','capabilities'],env=env,cwd=tmp_path,capture_output=True,text=True,encoding='utf8',timeout=30)
    assert p.returncode==0 and not json.loads(p.stdout)['requires_professional_package']
    p=subprocess.run([sys.executable,'-m','fintube_hx','field-import'],env=env,cwd=tmp_path,capture_output=True,timeout=30)
    assert p.returncode==2
    assert capabilities()['edition']=='community'


def test_no_private_module_imported():
    assert not any(n.startswith(('fintube_hx_pro','jiefeng_coil')) for n in sys.modules)
