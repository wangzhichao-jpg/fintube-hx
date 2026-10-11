"""Native Community XLSX: actual baseline data, snapshots, sources and limitations."""
import json
import math
import os
from pathlib import Path
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from .reporting import summary_record


def export_workbook(path,results):
    if not results:raise ValueError('No point records to export')
    wb=Workbook();summary=wb.active;summary.title='汇总';flat=[summary_record(r) for r in results]
    keys=list(dict.fromkeys(k for r in flat for k in r));summary.append(keys)
    for r in flat:summary.append([r.get(k) for k in keys])
    inputs=wb.create_sheet('输入快照');inputs.append(['case','part','input_json_utf8'])
    notes=wb.create_sheet('来源与限制');notes.append(['case','field','value']);notes.append(['ALL','verification','Community基础数值模型；未验证实际产品精度。空压降不是0。'])
    details=wb.create_sheet('轴向分段');cells=[]
    for r in results:
        name=r.get('case_name','failed')
        for label,key in ((name,'inputs'),(name+' [solved operating inputs]','operating_inputs')):
            if key in r:
                text=json.dumps(r[key],ensure_ascii=False,indent=2,allow_nan=False)
                for i in range(0,len(text),30000):inputs.append([label,i//30000+1,text[i:i+30000]])
        for key,value in r.get('provenance',{}).items():notes.append([name,key,value])
        for text in r.get('warnings',[]):notes.append([name,'warning',text])
        if 'error' in r:notes.append([name,'error',r['error']])
        for cell in r.get('segments',[]):cells.append({'case':name,**cell})
    keys=list(dict.fromkeys(k for c in cells for k in c)) or ['case','note'];details.append(keys)
    for c in cells:details.append([c.get(k) for k in keys])
    if not cells:details.append(['ALL','整体模型没有逐管/分段数据，不伪造。'])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(name='Microsoft YaHei',bold=True);c.fill=PatternFill('solid',fgColor='DCE8EE')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,float) and not math.isfinite(c.value):raise ValueError('Nonfinite Excel value')
                if isinstance(c.value,str):c.data_type='s';c.alignment=Alignment(wrap_text=True,vertical='top')
                elif isinstance(c.value,(int,float)) and not isinstance(c.value,bool):c.number_format='0.000000'
                if ws==summary and c.value is None and 'pressure_drop' in str(ws.cell(1,c.column).value):c.comment=Comment('未建模/不适用，不是零压降。','Fintube HX')
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(42,max(18,len(str(column[0].value))+2))
    inputs.column_dimensions['C'].width=100;notes.column_dimensions['C'].width=100
    target=Path(path);target.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix=target.name+'.',suffix='.xlsx',dir=target.parent);os.close(fd)
    try:wb.save(tmp);os.replace(tmp,target)
    except BaseException:
        if os.path.exists(tmp):os.unlink(tmp)
        raise
    finally:wb.close()
